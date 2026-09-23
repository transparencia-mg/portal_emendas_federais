"""
Prepara as planilhas do painel de Emendas Federais (não envia ao GitHub).

Fica na pasta scripts do repositório. Normalmente é chamado pelo atualizar.py
(na raiz), que também atualiza os links e envia tudo ao GitHub. Para rodar só esta etapa:

    python scripts/atualizar_bases.py

O que o script faz, nesta ordem:

 1. Lê upload/EMENDASFEDERAIS.xlsx (abas datas, pagamento e pagamentoRP), exclui a
    primeira linha e a primeira coluna de cada aba e grava:
        datas        -> datas.xlsx                 (raiz do repositório)
        pagamento    -> upload/pagamentos.xlsx
        pagamentoRP  -> upload/pagamentosrp.xlsx
 2. Lê upload/Dados Emendas - ATUALIZADO <data>.xlsx (a data no nome pode variar) e grava:
        1ª aba                          -> upload/dados_gerais_emendas.xlsx
        EXEC_TRANSFERÊNCIAS_ESPECIAIS   -> upload/execucao_transferencias_especiais_pix.xlsx
        PLANO_TRANSFERÊNCIAS_ESPECIAIS  -> upload/plano_execucao_emendas_pix.xlsx
 3. Confere se cada arquivo tem as colunas que o painel usa. Se faltar alguma, para
    ANTES de gravar e de apagar os originais.
 4. Apaga EMENDASFEDERAIS.xlsx e Dados Emendas - ATUALIZADO <data>.xlsx
    (quando chamado pelo atualizar.py, só depois que a atualização dos links também der certo).

Não é preciso rodar as duas partes juntas: se só um dos arquivos de origem estiver na
pasta, o script processa só ele.
"""
import re
import sys
import unicodedata
import warnings
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


# =====================================================
# CONFIGURAÇÃO
# =====================================================

BASE = Path(__file__).resolve().parent.parent   # raiz do repositório (este arquivo fica em scripts/)
PASTA_UPLOAD = BASE / "upload"

# Arquivos de origem (procurados na pasta upload e, se não estiverem lá, na raiz)
NOME_EMENDAS_FEDERAIS = "EMENDASFEDERAIS.xlsx"
PADRAO_DADOS_EMENDAS = "Dados Emendas - ATUALIZADO*.xlsx"   # a data no nome pode mudar

APAGAR_ORIGINAIS = True       # apaga os arquivos de origem ao final

# Segurança: se a primeira linha/coluna das abas de EMENDASFEDERAIS tiver dados,
# o script para (em vez de apagar informação). Troque para True para excluir mesmo assim.
EXCLUIR_MESMO_COM_DADOS = False

# Abas de EMENDASFEDERAIS.xlsx -> arquivo de saída
ABAS_EMENDAS_FEDERAIS = [
    ("datas",       BASE / "datas.xlsx"),
    ("pagamento",   PASTA_UPLOAD / "pagamentos.xlsx"),
    ("pagamentoRP", PASTA_UPLOAD / "pagamentosrp.xlsx"),
]

# Abas de "Dados Emendas - ATUALIZADO" -> arquivo de saída
# (None = primeira aba da planilha, qualquer que seja o nome)
ABAS_DADOS_EMENDAS = [
    (None,                             PASTA_UPLOAD / "dados_gerais_emendas.xlsx"),
    ("EXEC_TRANSFERÊNCIAS_ESPECIAIS",  PASTA_UPLOAD / "execucao_transferencias_especiais_pix.xlsx"),
    ("PLANO_TRANSFERÊNCIAS_ESPECIAIS", PASTA_UPLOAD / "plano_execucao_emendas_pix.xlsx"),
]

# Colunas que o painel (index.html) precisa em cada arquivo.
# Os nomes são comparados sem acento, sem maiúsculas e sem pontuação.
COLUNAS_OBRIGATORIAS = {
    "datas.xlsx": ["CODIGO SIAFI", "DATA INICIO", "DATA FIM"],
    "pagamentos.xlsx": ["CODIGO SIAFI", "NUMERO EMPENHO", "DATA EMPENHO", "CODIGO FONTE",
                        "VALOR EMPENHADO", "VALOR LIQUIDADO", "VALOR PAGO"],
    "pagamentosrp.xlsx": ["CODIGO SIAFI", "NUMERO EMPENHO", "DATA EMPENHO", "CODIGO FONTE",
                          "VALOR EMPENHADO", "VALOR CANCELADO RPP", "VALOR CANCELADO RPNP",
                          "VALOR PAGO RPP", "VALOR PAGO RPNP"],
    "dados_gerais_emendas.xlsx": ["CÓDIGO SIAFI", "TIPO DE INSTRUMENTO JURÍDICO", "VALOR REPASSADO"],
    "execucao_transferencias_especiais_pix.xlsx": ["Nº SIAFI", "Nº Empenho"],
    "plano_execucao_emendas_pix.xlsx": ["codigo emenda e parlamentar"],
}
# Colunas desejáveis: se faltarem, só aparece um aviso
# (a coluna Link de dados_gerais_emendas.xlsx é preenchida depois, pelo atualizar_links.py)
COLUNAS_RECOMENDADAS = {
    "plano_execucao_emendas_pix.xlsx": ["Nº SIAFI"],
}


# =====================================================
# FUNÇÕES AUXILIARES
# =====================================================

class ErroProcessamento(Exception):
    pass


def normalizar(texto):
    """'Nº SIAFI' e 'n siafi' viram 'N SIAFI'; 'EXEC_TRANSFERÊNCIAS' vira 'EXEC TRANSFERENCIAS'."""
    s = unicodedata.normalize("NFD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^A-Za-z0-9]+", " ", s).upper().split())


def vazio(v):
    return v is None or (isinstance(v, str) and not v.strip())


def limpar_valor(v):
    """Números inteiros gravados como decimal (9241504.0) voltam a ser inteiros,
    para o SIAFI não virar '9241504,0' no CSV gerado pelo GitHub."""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def ler_tabela(ws, excluir_primeira_linha=False, excluir_primeira_coluna=False):
    """Lê a aba como lista de linhas, já sem linhas e colunas totalmente vazias."""
    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    if not linhas:
        raise ErroProcessamento(f"A aba '{ws.title}' está vazia.")

    if excluir_primeira_linha:
        if any(not vazio(v) for v in linhas[0]) and not EXCLUIR_MESMO_COM_DADOS:
            raise ErroProcessamento(
                f"A primeira linha da aba '{ws.title}' tem dados: {[v for v in linhas[0] if not vazio(v)][:5]}. "
                "Confira a planilha (ou use EXCLUIR_MESMO_COM_DADOS = True).")
        linhas = linhas[1:]
    if excluir_primeira_coluna:
        preenchidas = [r[0] for r in linhas if r and not vazio(r[0])]
        if preenchidas and not EXCLUIR_MESMO_COM_DADOS:
            raise ErroProcessamento(
                f"A primeira coluna da aba '{ws.title}' tem dados: {preenchidas[:5]}. "
                "Confira a planilha (ou use EXCLUIR_MESMO_COM_DADOS = True).")
        linhas = [r[1:] for r in linhas]

    # Remove linhas totalmente vazias (inclusive as de cima, antes do cabeçalho)
    linhas = [r for r in linhas if any(not vazio(v) for v in r)]
    if not linhas:
        raise ErroProcessamento(f"A aba '{ws.title}' não tem dados.")

    # Remove colunas totalmente vazias (inclusive as de formatação, à direita)
    largura = max(len(r) for r in linhas)
    linhas = [r + [None] * (largura - len(r)) for r in linhas]
    manter = [c for c in range(largura) if any(not vazio(r[c]) for r in linhas)]
    linhas = [[r[c] for c in manter] for r in linhas]

    cabecalho = [str(v).strip() if not vazio(v) else "" for v in linhas[0]]
    sem_nome = [i + 1 for i, c in enumerate(cabecalho) if not c]
    if sem_nome:
        raise ErroProcessamento(f"A aba '{ws.title}' tem coluna(s) com dados mas sem cabeçalho (posição {sem_nome}).")
    repetidas = sorted({c for c in cabecalho if cabecalho.count(c) > 1})
    if repetidas:
        raise ErroProcessamento(f"A aba '{ws.title}' tem colunas com nome repetido: {repetidas}.")

    dados = [[limpar_valor(v) for v in r] for r in linhas[1:]]
    return cabecalho, dados


def conferir_colunas(nome_arquivo, cabecalho):
    presentes = {normalizar(c) for c in cabecalho}
    faltando = [c for c in COLUNAS_OBRIGATORIAS.get(nome_arquivo, []) if normalizar(c) not in presentes]
    if faltando:
        raise ErroProcessamento(
            f"{nome_arquivo}: faltam colunas usadas pelo painel: {faltando}.\n"
            f"    Colunas encontradas: {cabecalho}")
    for c in COLUNAS_RECOMENDADAS.get(nome_arquivo, []):
        if normalizar(c) not in presentes:
            print(f"    AVISO: {nome_arquivo} não tem a coluna '{c}'.")


def gravar_xlsx(caminho, titulo_aba, cabecalho, dados):
    wb = Workbook()
    ws = wb.active
    ws.title = (titulo_aba or "Planilha1")[:31]
    ws.append(cabecalho)
    for linha in dados:
        ws.append(linha)
    # Datas com formato de data (sem hora), para leitura correta no painel e no CSV
    for col in ws.iter_cols(min_row=2):
        for cel in col:
            if isinstance(cel.value, datetime) and cel.value.time() == datetime.min.time():
                cel.number_format = "DD/MM/YYYY"
            elif isinstance(cel.value, date) and not isinstance(cel.value, datetime):
                cel.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    temporario = caminho.with_name(caminho.stem + ".tmp.xlsx")
    wb.save(temporario)
    temporario.replace(caminho)   # só substitui o arquivo antigo depois de gravar o novo inteiro


def achar_aba(wb, nome, arquivo):
    if nome is None:
        return wb.worksheets[0]
    alvo = normalizar(nome)
    for ws in wb.worksheets:
        if normalizar(ws.title) == alvo:
            return ws
    raise ErroProcessamento(f"Aba '{nome}' não encontrada em {arquivo}. Abas existentes: {wb.sheetnames}")


def achar_arquivo(padrao):
    """Procura na pasta upload e depois na raiz. Se houver vários, usa o mais recente."""
    for pasta in (PASTA_UPLOAD, BASE):
        achados = [p for p in pasta.glob(padrao) if not p.name.startswith("~$")]
        if achados:
            achados.sort(key=lambda p: p.stat().st_mtime, reverse=True)
            if len(achados) > 1:
                print(f"  AVISO: {len(achados)} arquivos '{padrao}'. Usando o mais recente: {achados[0].name}")
            return achados[0]
    return None


def processar(origem, abas, excluir_linha_coluna):
    """Gera os arquivos de saída a partir das abas de 'origem'. Retorna a lista gravada."""
    print(f"\nLendo {origem.name}...")
    try:
        wb = load_workbook(origem, data_only=True, read_only=True)
    except PermissionError:
        raise ErroProcessamento(f"Não foi possível abrir {origem.name}. Feche o arquivo no Excel e rode de novo.")

    preparados = []
    try:
        for nome_aba, destino in abas:
            ws = achar_aba(wb, nome_aba, origem.name)
            cabecalho, dados = ler_tabela(ws, excluir_linha_coluna, excluir_linha_coluna)
            conferir_colunas(destino.name, cabecalho)
            preparados.append((ws.title, destino, cabecalho, dados))
    finally:
        wb.close()

    # Só grava depois de conferir todas as abas deste arquivo
    for titulo, destino, cabecalho, dados in preparados:
        try:
            gravar_xlsx(destino, titulo, cabecalho, dados)
        except PermissionError:
            raise ErroProcessamento(f"Não foi possível gravar {destino.name}. Feche o arquivo no Excel e rode de novo.")
        local = destino.relative_to(BASE)
        n = f"{len(dados):,}".replace(",", ".")
        print(f"  aba '{titulo}' -> {local}  ({n} linhas, {len(cabecalho)} colunas)")
    return [p[1] for p in preparados]


# =====================================================
# EXECUÇÃO
# =====================================================

def localizar_origens():
    """Retorna (EMENDASFEDERAIS, Dados Emendas - ATUALIZADO); None para o que não existir."""
    if not PASTA_UPLOAD.is_dir():
        raise ErroProcessamento(f"Pasta upload não encontrada em {BASE}.")
    arq_federais = achar_arquivo(NOME_EMENDAS_FEDERAIS)
    arq_dados = achar_arquivo(PADRAO_DADOS_EMENDAS)
    if not arq_federais:
        print(f"  {NOME_EMENDAS_FEDERAIS} não encontrado: datas/pagamentos/pagamentosrp não serão atualizados.")
    if not arq_dados:
        print("  'Dados Emendas - ATUALIZADO' não encontrado: dados gerais/execução PIX/plano não serão atualizados.")
    return arq_federais, arq_dados


def executar():
    """Gera as planilhas. Retorna (arquivos_gerados, arquivos_de_origem).
    Não apaga os originais: use apagar_originais() depois que tudo der certo."""
    arq_federais, arq_dados = localizar_origens()
    gerados = []
    if arq_federais:
        gerados += processar(arq_federais, ABAS_EMENDAS_FEDERAIS, excluir_linha_coluna=True)
    if arq_dados:
        gerados += processar(arq_dados, ABAS_DADOS_EMENDAS, excluir_linha_coluna=False)
    return gerados, [a for a in (arq_federais, arq_dados) if a]


def apagar_originais(origens):
    """Apaga os arquivos de origem. Retorna False se algum não pôde ser apagado."""
    ok = True
    for origem in origens:
        try:
            origem.unlink()
            print(f"  apagado: {origem.relative_to(BASE)}")
        except FileNotFoundError:
            pass
        except PermissionError:
            print(f"  ERRO: não foi possível apagar {origem.name} (está aberto no Excel?).")
            ok = False
    return ok


def main():
    print(f"Repositório: {BASE}")
    try:
        gerados, origens = executar()
    except ErroProcessamento as erro:
        print(f"\nERRO: {erro}")
        sys.exit("Processo interrompido. Os arquivos de origem NÃO foram apagados.")
    if not origens:
        print("Nenhum arquivo de origem na pasta upload. Nada a fazer.")
        return
    if APAGAR_ORIGINAIS:
        print("\nApagando arquivos de origem...")
        apagar_originais(origens)
    print("\nConcluído (sem envio ao GitHub; para enviar, use: python atualizar.py).")


if __name__ == "__main__":
    main()
